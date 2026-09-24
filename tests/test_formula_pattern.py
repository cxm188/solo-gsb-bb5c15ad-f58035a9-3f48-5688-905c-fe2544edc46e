"""同列公式模板归一化与相邻区段偏离检查的核心逻辑测试。"""
import io

import pytest
from openpyxl import Workbook, load_workbook

from app.formula_pattern import (
    MIN_SEGMENT_CELLS,
    analyze_deviations,
    normalize_formula,
)
from app.graph import WorkbookGraph


def load_graph(wb: Workbook) -> WorkbookGraph:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wbf = load_workbook(buf, data_only=False)
    buf.seek(0)
    wbv = load_workbook(buf, data_only=True)
    return WorkbookGraph.from_workbooks(wbf, wbv)


# ------------------------------------------------------------ 模板归一化
def test_relative_and_absolute_share_template():
    # 需求示例：C2 的 =A2+$B$1 与 C3 的 =A3+$B$1 属于同一模板
    t2 = normalize_formula("=A2+$B$1", owner_col=3, owner_row=2)
    t3 = normalize_formula("=A3+$B$1", owner_col=3, owner_row=3)
    assert t2 == t3 == "=R[0]C[-2]+R1C2"


def test_fully_relative_range_shifts():
    t2 = normalize_formula("=SUM(A2:B3)", 5, 2)
    t3 = normalize_formula("=SUM(A3:B4)", 5, 3)
    assert t2 == t3 == "=SUM(R[0]C[-4]:R[1]C[-3])"


def test_absolute_dimensions_stay_fixed():
    # $A2：列绝对（写死 C1）、行相对（同行 R[0]），向下复制模板一致
    t2 = normalize_formula("=$A2", 3, 2)
    t3 = normalize_formula("=$A3", 3, 3)
    assert t2 == t3 == "=R[0]C1"
    # B$2：列相对（左移一格 C[-1]）、行绝对（写死 R2）；向下复制时原文不变，
    # 两格归一到同一模板
    assert normalize_formula("=B$2", 3, 2) == normalize_formula("=B$2", 3, 3) \
        == "=R2C[-1]"
    # 绝对行被改成随行变化（B$1 vs B$2）时模板不同
    assert normalize_formula("=B$1", 3, 2) != normalize_formula("=B$2", 3, 3)
    # 绝对列被改成随列变化时模板不同
    assert normalize_formula("=$B1", 3, 2) != normalize_formula("=$A1", 3, 2)


def test_cross_sheet_qualifier_preserved():
    # 跨表部分：计算!B2 相对平移；参数!$B$1 绝对行列固定
    t2 = normalize_formula("=计算!B2+参数!$B$1", 3, 2)
    t3 = normalize_formula("=计算!B3+参数!$B$1", 3, 3)
    assert t2 == t3 == "=计算!R[0]C[-1]+参数!R1C2"


def test_quoted_sheet_name_kept():
    t = normalize_formula("='月报 2026'!B2", 3, 2)
    assert t == "='月报 2026'!R[0]C[-1]"


def test_whole_column_and_row_refs():
    t2 = normalize_formula("=SUM(D:D)+SUM(2:2)", 3, 2)
    t3 = normalize_formula("=SUM(D:D)+SUM(3:3)", 3, 3)
    assert t2 == "=SUM(C[1]:C[1])+SUM(R[0]:R[0])"
    # 整行相对引用随行平移，模板一致；绝对整行则保持
    assert t2 == t3
    assert normalize_formula("=SUM($D:$D)", 3, 2) == "=SUM(C4:C4)"


def test_strings_and_non_reference_text_verbatim():
    # 字符串里的 "B2" 不参与归一；函数名 / 数字 / 运算符原样保留
    t = normalize_formula('=IF(A2="B2",$B$1,"A9")', 3, 2)
    assert t == '=IF(R[0]C[-2]="B2",R1C2,"A9")'


def test_defined_names_preserved_verbatim():
    # 定义名称（含跨表限定）原样保留
    t2 = normalize_formula("=单价*A2", 3, 2)
    t3 = normalize_formula("=单价*A3", 3, 3)
    assert t2 == t3 == "=单价*R[0]C[-2]"
    assert normalize_formula("=计算!系数", 3, 2) == "=计算!系数"


def test_error_literal_preserved():
    assert normalize_formula("=#REF!+A2", 3, 2) == "=#REF!+R[0]C[-2]"


def test_leading_equals_optional():
    assert normalize_formula("A2+$B$1", 3, 2) == normalize_formula("=A2+$B$1", 3, 2)


# ------------------------------------------------------------ 区段分析
def _segments_map(res):
    return {s["segment"]: s for s in res["segments"]}


def test_strict_majority_flags_other_formulas():
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    for r in range(2, 7):                 # C2:C6，主流 4 格
        ws[f"C{r}"] = f"=A{r}+$B$1"
    ws["C5"] = "=A5+999"                  # 1 格偏离
    g = load_graph(wb)
    res = g.formula_deviations()
    seg = _segments_map(res)["C2:C6"]
    assert seg["sheet"] == "S1"
    assert seg["column"] == "C"
    assert seg["segment_start"] == "C2" and seg["segment_end"] == "C6"
    assert seg["segment_rows"] == 5
    assert seg["majority_count"] == 4
    assert seg["majority_example_formula"] == "=A2+$B$1"
    assert seg["majority_example_cell"] == "S1!C2"
    assert seg["deviation_count"] == 1
    odd = seg["deviations"][0]
    assert odd["cell"] == "S1!C5" and odd["coord"] == "C5"
    assert odd["display"] == "S1!C5"
    assert odd["formula"] == "=A5+999"
    assert odd["pattern"] == "=R[0]C[-2]+999"
    # pattern_counts 覆盖全部模板
    assert {p["pattern"]: p["count"] for p in seg["pattern_counts"]} == {
        "=R[0]C[-2]+R1C2": 4, "=R[0]C[-2]+999": 1}
    assert res["total"] == 1
    assert res["deviation_cell_count"] == 1
    assert res["anomalous_segment_count"] == 1


def test_tie_is_not_reported():
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws["B2"], ws["B3"] = "=A2", "=A3"
    ws["B4"], ws["B5"] = "=A4*2", "=A5*2"   # 2:2 平票
    ws["C2"], ws["C3"] = "=A2", "=A3"
    ws["C4"] = "=X4"                        # 3 格区段内 2:1 也不足多数？2 > 3/2
    g = load_graph(wb)
    res = g.formula_deviations()
    # B2:B5 平票不报；C2:C4 主流 2 格严格过半（2*2>3）报 1 格
    assert "B2:B5" not in _segments_map(res)
    cseg = _segments_map(res)["C2:C4"]
    assert [d["coord"] for d in cseg["deviations"]] == ["C4"]


def test_three_way_split_no_majority_not_reported():
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws["B2"], ws["B3"], ws["B4"], ws["B5"], ws["B6"] = (
        "=A2", "=A3", "=A4*2", "=A5*3", "=A6*4")  # 计数 2:1:1:1，最高仅 2 未过半
    g = load_graph(wb)
    res = g.formula_deviations()
    assert res["total"] == 0
    assert res["anomalous_segment_count"] == 0
    # 区段仍计入「参与检查」
    assert res["checked_segment_count"] == 1


def test_uniform_segment_not_reported():
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    for r in range(2, 6):
        ws[f"B{r}"] = f"=A{r}+1"
    g = load_graph(wb)
    res = g.formula_deviations()
    assert res["segments"] == []
    assert res["checked_segment_count"] == 1
    assert res["anomalous_segment_count"] == 0


def test_segments_shorter_than_three_skipped():
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws["B2"] = "=A2"
    ws["B3"] = "=Z99"          # 2 格：不足 3 格，即便模板不一致也不报
    g = load_graph(wb)
    res = g.formula_deviations()
    assert res["segments"] == []
    assert res["checked_segment_count"] == 0


def test_blank_row_breaks_segment():
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws["B2"], ws["B3"] = "=A2", "=A3"
    # B4 空
    ws["B5"], ws["B6"], ws["B7"] = "=A5", "=Z6", "=A7"  # 3 格区段，2 格主流
    g = load_graph(wb)
    res = g.formula_deviations()
    # 只有 B5:B7 一段参与检查，B6 报偏离
    assert res["checked_segment_count"] == 1
    seg = _segments_map(res)["B5:B7"]
    assert [d["coord"] for d in seg["deviations"]] == ["B6"]


def test_different_columns_are_independent():
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    for r in range(2, 6):
        ws[f"B{r}"] = f"=A{r}+1"
    for r in range(2, 6):
        ws[f"C{r}"] = f"=B{r}*2"
    ws["C4"] = "=B4/2"          # C 列 3:1 多数 => 报
    g = load_graph(wb)
    res = g.formula_deviations()
    assert list(_segments_map(res)) == ["C2:C5"]


def test_sheet_filter_and_404():
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    for r in range(2, 5):
        ws[f"B{r}"] = f"=A{r}+1"
    ws2 = wb.create_sheet("S2")
    for r in range(2, 5):
        ws2[f"B{r}"] = f"=A{r}+1"
    ws2["B3"] = "=Z99"
    g = load_graph(wb)
    all_res = g.formula_deviations()
    assert all_res["total"] == 1 and all_res["segments"][0]["sheet"] == "S2"
    assert all_res["deviation_cell_count"] == 1
    assert all_res["total_deviation_cell_count"] == 1
    f1 = g.formula_deviations(sheet="S1")
    assert f1["segments"] == [] and f1["total"] == 0
    # 可见偏离为 0，但全工作簿计数仍保留
    assert f1["deviation_cell_count"] == 0
    assert f1["total_deviation_cell_count"] == 1
    assert f1["anomalous_segment_count"] == 1
    # sheets 汇总仍覆盖整本工作簿
    assert {s["name"] for s in f1["sheets"]} == {"S1", "S2"}
    f2 = g.formula_deviations(sheet="S2")
    assert len(f2["segments"]) == 1
    assert f2["deviation_cell_count"] == 1


def test_min_segment_constant_is_three():
    assert MIN_SEGMENT_CELLS == 3


def test_segments_sorted_by_sheet_order_then_column_then_row():
    wb = Workbook()
    wb.active.title = "AA"
    wb.create_sheet("ZZ")
    for sheet, col in (("ZZ", "D"), ("AA", "B"), ("AA", "C")):
        ws = wb[sheet]
        for r in range(2, 5):
            ws[f"{col}{r}"] = f"=A{r}+1"
        ws[f"{col}3"] = "=Z99"
    g = load_graph(wb)
    res = g.formula_deviations()
    keys = [(s["sheet"], s["column"], s["segment_start"]) for s in res["segments"]]
    assert keys == [("AA", "B", "B2"), ("AA", "C", "C2"), ("ZZ", "D", "D2")]

    # 同一表内多个区段按起始行排序
    wb2 = Workbook()
    w2 = wb2.active
    w2.title = "S1"
    for col in ("B", "D"):
        for r in range(20, 24):
            w2[f"{col}{r}"] = f"=A{r}+1"
        w2[f"{col}21"] = "=Z99"
    g2 = load_graph(wb2)
    res2 = g2.formula_deviations()
    assert [s["segment"] for s in res2["segments"]] == ["B20:B23", "D20:D23"]


def test_sheet_stats_counts():
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    for r in range(2, 5):
        ws[f"B{r}"] = f"=A{r}+1"      # 1 段，无异常
    ws2 = wb.create_sheet("S2")
    for r in range(2, 7):
        ws2[f"C{r}"] = f"=A{r}+1"
    ws2["C3"] = "=Z99"                # 1 段异常，1 格偏离
    g = load_graph(wb)
    res = g.formula_deviations()
    stats = {s["name"]: s for s in res["sheets"]}
    assert stats["S1"] == {"name": "S1", "checked_segments": 1,
                           "anomalous_segments": 0, "deviation_cells": 0}
    assert stats["S2"]["anomalous_segments"] == 1
    assert stats["S2"]["deviation_cells"] == 1


def test_analyze_deviations_empty():
    res = analyze_deviations({}, [])
    assert res["segments"] == [] and res["total"] == 0
    assert res["sheets"] == [] and res["formula_cell_count"] == 0
